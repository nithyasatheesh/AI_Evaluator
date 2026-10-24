"""Generate per-submission reports and a consolidated participant feedback register."""
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REGISTER_NAME = "Participant_Feedback_Register.xlsx"


def _participant_name(source_file: str) -> str:
    path = Path(source_file)
    # Prefer an explicit Participant_NNN folder or filename.
    candidates = [path.stem, *reversed(path.parts[:-1])]
    for item in candidates:
        match = re.search(r"(participant[_ -]*\d+)", item, re.IGNORECASE)
        if match:
            return re.sub(r"[\s-]+", "_", match.group(1)).replace("__", "_")
    return path.parent.name or path.stem


def _feedback_text(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "\n".join(str(item) for item in value if item is not None)
    return str(value)


def _format_worksheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value or "")
        max_len = max((len(str(c.value or "")) for c in column_cells), default=12)
        if header in {"Area of Strength", "Area for Improvement", "Overall Feedback"}:
            width = 48
        elif header == "Submission File":
            width = 42
        else:
            width = min(32, max(12, max_len + 2))
        ws.column_dimensions[letter].width = width
        for cell in column_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _criterion_column(name, max_marks):
    return f"{name} ({max_marks})"


def _register_row(result, source_file):
    row = {
        "Participant Name": _participant_name(source_file),
        "Submission File": source_file.replace("\\", "/"),
        "Project": result.get("project_id", "Loan_Default"),
        "Milestone": result.get("milestone_id", ""),
    }
    for criterion in result.get("criteria", []):
        name = criterion.get("name", "Unnamed Criterion")
        max_marks = criterion.get("max_marks", "")
        row[_criterion_column(name, max_marks)] = criterion.get("marks_awarded", 0)
    row["Total Score"] = result.get("total_marks", 0)
    row["Maximum Score"] = result.get("maximum_marks", sum(
        c.get("max_marks", 0) for c in result.get("criteria", [])
    ))
    row["Area of Strength"] = _feedback_text(result.get("strengths", []))
    row["Area for Improvement"] = _feedback_text(result.get("areas_for_improvement", []))
    row["Overall Feedback"] = _feedback_text(result.get("overall_feedback", ""))
    return row


def _write_consolidated_register(result, source_file, output_dir):
    out = Path(output_dir)
    register_path = out / REGISTER_NAME
    new_row = _register_row(result, source_file)

    # Keep existing rows in the current run, updating the same submission if rerun.
    if register_path.exists():
        try:
            existing = pd.read_excel(register_path, sheet_name="Feedback Register").fillna("")
            rows = existing.to_dict(orient="records")
        except (ValueError, OSError, ImportError):
            rows = []
    else:
        rows = []

    key = (str(new_row["Submission File"]), str(new_row["Milestone"]))
    rows = [
        row for row in rows
        if (str(row.get("Submission File", "")), str(row.get("Milestone", ""))) != key
    ]
    rows.append(new_row)

    # Stable identifying columns first, score columns next, feedback columns last.
    base = ["Participant Name", "Submission File", "Project", "Milestone"]
    feedback = ["Area of Strength", "Area for Improvement", "Overall Feedback"]
    all_keys = list(dict.fromkeys(key for row in rows for key in row.keys()))
    score_cols = [
        key for key in all_keys
        if key not in base + feedback + ["Total Score", "Maximum Score"]
    ]
    ordered = base + score_cols + ["Total Score", "Maximum Score"] + feedback
    frame = pd.DataFrame(rows).reindex(columns=ordered).fillna("")
    # Keep one record per participant submission and milestone.
    frame = frame.drop_duplicates(subset=["Submission File", "Milestone"], keep="last")
    with pd.ExcelWriter(register_path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Feedback Register")
        _format_worksheet(writer.book["Feedback Register"])
    return register_path


def write_reports(result, source_file, output_dir):
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    stem = Path(source_file).stem
    xlsx = out / f"{stem}_evaluation.xlsx"
    js = out / f"{stem}_evaluation.json"

    payload = dict(result)
    payload["source_file"] = source_file
    payload["human_review_required"] = True
    js.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

    criteria_rows = [{
        "Project": result.get("project_id"),
        "Milestone": result.get("milestone_id"),
        "Criterion": c.get("name"),
        "Maximum Marks": c.get("max_marks"),
        "Marks Awarded": c.get("marks_awarded"),
        "Evidence": c.get("evidence", ""),
        "Feedback / Improvement": c.get("feedback", ""),
    } for c in result.get("criteria", [])]

    summary_rows = [
        ["Project", result.get("project_id")],
        ["Milestone", result.get("milestone_id")],
        ["Total Marks", result.get("total_marks")],
        ["Maximum Marks", result.get("maximum_marks")],
        ["Human Review Required", True],
        ["Overall Feedback", _feedback_text(result.get("overall_feedback", ""))],
        ["Strengths", _feedback_text(result.get("strengths", []))],
        ["Areas for Improvement", _feedback_text(result.get("areas_for_improvement", []))],
        ["Source File", source_file],
    ]
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        pd.DataFrame(criteria_rows).to_excel(writer, index=False, sheet_name="Criteria")
        pd.DataFrame(summary_rows, columns=["Field", "Value"]).to_excel(
            writer, index=False, sheet_name="Summary"
        )
        for ws in writer.book.worksheets:
            _format_worksheet(ws)

    register_path = _write_consolidated_register(result, source_file, out)
    return xlsx, js
